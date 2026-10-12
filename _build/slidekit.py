"""slidekit.py -- shared slide builder for the CSS SKL 342 lab decks.

Each lab has a small script (labN_slides.py) that describes its slides with the
helpers below; this module owns the look (16:9, UW purple) so all decks match.

Inline code: wrap it in backticks inside any bullet or step, e.g. "use `const&`".
Requires python-pptx (pip install python-pptx).
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------- palette
PURPLE = RGBColor(0x4B, 0x2E, 0x83)
DARK = RGBColor(0x1E, 0x14, 0x38)
GOLD = RGBColor(0xC5, 0xB4, 0xE3)       # light lavender accent on dark slides
TEXT = RGBColor(0x1D, 0x23, 0x30)
MUTED = RGBColor(0x5B, 0x64, 0x75)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE = RGBColor(0xF3, 0xF1, 0xF8)
BORDER = RGBColor(0xE2, 0xE6, 0xEE)
CODE_BG = RGBColor(0x1E, 0x1E, 0x2E)
CODE_TEXT = RGBColor(0xE6, 0xE6, 0xEF)
CODE_COMMENT = RGBColor(0x9A, 0xA4, 0xB5)
LECTURE = RGBColor(0x2F, 0x6F, 0xDF)
HANDSON = RGBColor(0x1F, 0x9D, 0x55)
SOLUTION = RGBColor(0xB7, 0x79, 0x1F)
KIND_COLOR = {"Lecture": LECTURE, "Hands-on": HANDSON, "Solutions": SOLUTION, "Quiz": PURPLE}

BODY_FONT = "Calibri"
HEAD_FONT = "Calibri"
CODE_FONT = "Consolas"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)


class Deck:
    def __init__(self, lab_label):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.lab_label = lab_label

    # ------------------------------------------------------------ primitives
    def _slide(self, dark=False):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = DARK if dark else WHITE
        if not dark:
            self._footer(s)
        return s

    def _footer(self, s):
        tb = self._text(s, Inches(0.6), Inches(7.0), Inches(8), Inches(0.3), "")
        self._runs(tb.text_frame.paragraphs[0], f"CSS SKL 342 · {self.lab_label}", 11, MUTED)
        num = self._text(s, Inches(11.7), Inches(7.0), Inches(1.0), Inches(0.3), "")
        p = num.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        self._runs(p, str(len(self.prs.slides)), 11, MUTED)

    def _text(self, s, x, y, w, h, text, size=16, color=TEXT, bold=False, font=BODY_FONT,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(margin)
        tf.vertical_anchor = anchor
        if text:
            p = tf.paragraphs[0]
            p.alignment = align
            self._runs(p, text, size, color, bold, font)
        return tb

    @staticmethod
    def _runs(p, text, size, color, bold=False, font=BODY_FONT, code_color=None):
        """Add text to paragraph p; `backticked` spans are set in the code font."""
        parts = text.split("`")
        for i, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            r.font.size = Pt(size)
            r.font.bold = bold
            if i % 2 == 1:
                r.font.name = CODE_FONT
                r.font.color.rgb = code_color or PURPLE
            else:
                r.font.name = font
                r.font.color.rgb = color

    @staticmethod
    def _bullet(p, level=0, char="•"):
        pPr = p._p.get_or_add_pPr()
        indent = 0.28 + 0.3 * level
        pPr.set("marL", str(int(Inches(indent))))
        pPr.set("indent", str(int(-Inches(0.25))))
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", "5B6475")
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", char)

    def _box(self, s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
        b = s.shapes.add_shape(shape, x, y, w, h)
        b.fill.solid()
        b.fill.fore_color.rgb = fill
        if line is None:
            b.line.fill.background()
        else:
            b.line.color.rgb = line
            b.line.width = Pt(1)
        b.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            b.adjustments[0] = radius
        return b

    def _title(self, s, title, kind=None):
        x = MARGIN
        if kind:
            pill = self._box(s, MARGIN, Inches(0.5), Inches(1.45), Inches(0.36), KIND_COLOR[kind], radius=0.5)
            tf = pill.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            self._runs(p, kind.upper(), 12, WHITE, bold=True)
            x = MARGIN + Inches(1.65)
        self._text(s, x, Inches(0.38), W - x - MARGIN, Inches(0.7), title, 32, TEXT, bold=True,
                   font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)

    def _bullets(self, s, x, y, w, h, items, size=18):
        """items: str (level 0) or (str, level)."""
        tb = self._text(s, x, y, w, h, "")
        tf = tb.text_frame
        first = True
        for it in items:
            text, level = (it, 0) if isinstance(it, str) else it
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            self._bullet(p, level, "•" if level == 0 else "–")
            p.space_after = Pt(8 if level == 0 else 4)
            self._runs(p, text, size if level == 0 else size - 2, TEXT if level == 0 else MUTED)
        return tb

    def _code(self, s, x, y, w, h, code, size=None):
        lines = code.strip("\n").split("\n")
        if size is None:
            size = 16 if len(lines) <= 9 else 14 if len(lines) <= 14 else 12
        # shrink so the longest line fits (Consolas is ~0.55 em wide), then fit the panel to the text
        longest = max(len(l) for l in lines)
        fit = (w - Inches(0.5)) / Inches(1) * 72 / (0.55 * max(longest, 1))
        size = max(10, min(size, int(fit)))
        need = Inches(len(lines) * size * 1.2 / 72 + 0.45)
        h = min(h, max(need, Inches(1.0)))
        panel = self._box(s, x, y, w, h, CODE_BG, radius=0.03)
        tb = self._text(s, x + Inches(0.25), y + Inches(0.2), w - Inches(0.5), h - Inches(0.4), "")
        tf = tb.text_frame
        tf.word_wrap = False
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(0)
            code_part, _, comment = line.partition("//")
            if code_part:
                r = p.add_run()
                r.text = code_part
                r.font.name, r.font.size, r.font.color.rgb = CODE_FONT, Pt(size), CODE_TEXT
            if _:
                r = p.add_run()
                r.text = "//" + comment
                r.font.name, r.font.size, r.font.color.rgb = CODE_FONT, Pt(size), CODE_COMMENT
            if not line:
                r = p.add_run()
                r.text = " "
                r.font.name, r.font.size = CODE_FONT, Pt(size)
        return panel

    def _callout(self, s, x, y, w, h, head, body, tint=SURFACE, head_color=PURPLE, size=16):
        self._box(s, x, y, w, h, tint, radius=0.08)
        tb = self._text(s, x + Inches(0.25), y + Inches(0.15), w - Inches(0.5), h - Inches(0.3), "",
                        anchor=MSO_ANCHOR.MIDDLE)
        tf = tb.text_frame
        p = tf.paragraphs[0]
        self._runs(p, head, size, head_color, bold=True)
        p2 = tf.add_paragraph()
        p2.space_before = Pt(4)
        self._runs(p2, body, size, TEXT)

    def _circle_num(self, s, x, y, d, label, fill=PURPLE, color=WHITE, size=20):
        c = self._box(s, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        self._runs(p, label, size, color, bold=True)
        return c

    # ------------------------------------------------------------ slide types
    def title(self, eyebrow, title, subtitle, notes=""):
        s = self._slide(dark=True)
        self._box(s, Inches(9.3), Inches(-1.2), Inches(5.2), Inches(5.2), PURPLE, shape=MSO_SHAPE.OVAL)
        self._text(s, MARGIN, Inches(2.2), Inches(10), Inches(0.5), eyebrow.upper(), 16, GOLD, bold=True)
        size = 48 if len(title) <= 34 else 40          # longer titles shrink to stay on one line
        self._text(s, MARGIN, Inches(2.75), Inches(10.5), Inches(1.4), title, size, WHITE, bold=True, font=HEAD_FONT)
        self._text(s, MARGIN, Inches(4.25), Inches(9.5), Inches(1.2), subtitle, 20, GOLD)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def agenda(self, parts, notes=""):
        """parts: list of (num, title, [(kind, time, minutes, label)])."""
        s = self._slide()
        self._title(s, "Today's plan")
        n = len(parts)
        gap = Inches(0.3)
        cw = (W - 2 * MARGIN - gap * (n - 1)) / n
        y = Inches(1.45)
        for i, (num, title, rows) in enumerate(parts):
            x = MARGIN + i * (cw + gap)
            self._box(s, x, y, cw, Inches(5.25), SURFACE, radius=0.05)
            self._circle_num(s, x + Inches(0.25), y + Inches(0.25), Inches(0.6), str(num))
            self._text(s, x + Inches(0.25), y + Inches(1.0), cw - Inches(0.5), Inches(1.1), title, 19, TEXT, bold=True)
            ry = y + Inches(2.2)
            for kind, time, mins, label in rows:
                pill = self._box(s, x + Inches(0.25), ry, Inches(1.25), Inches(0.32), KIND_COLOR[kind], radius=0.5)
                tf = pill.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                self._runs(p, kind.upper(), 10, WHITE, bold=True)
                self._text(s, x + Inches(1.6), ry - Inches(0.02), cw - Inches(1.8), Inches(0.35),
                           f"{time} · {mins}", 13, MUTED)
                self._text(s, x + Inches(0.25), ry + Inches(0.38), cw - Inches(0.5), Inches(0.5), label, 14, TEXT)
                ry += Inches(1.0)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def divider(self, num, title, goals, notes="", why=None):
        """why: one line on why the topic matters (the page's intro is in notes)."""
        s = self._slide(dark=True)
        top = Inches(0.95) if why else Inches(1.6)       # with a why line, everything moves up
        self._circle_num(s, MARGIN, top, Inches(1.1), str(num), fill=PURPLE, size=40)
        y = top + Inches(1.4)
        self._text(s, MARGIN, y, Inches(12), Inches(1.0), title, 42, WHITE, bold=True, font=HEAD_FONT)
        y += Inches(1.15)
        if why:
            tb = self._text(s, MARGIN, y, Inches(12), Inches(0.5), "")
            self._runs(tb.text_frame.paragraphs[0], why, 20, GOLD, code_color=WHITE)
            y += Inches(0.85)
        self._text(s, MARGIN, y, Inches(12), Inches(0.4), "AFTER THIS PART YOU CAN", 14, GOLD, bold=True)
        tb = self._text(s, MARGIN, y + Inches(0.45), Inches(11.5), Inches(2.4), "")
        tf = tb.text_frame
        for i, g in enumerate(goals):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            self._bullet(p, 0, "›")
            p.space_after = Pt(6)
            self._runs(p, g, 20, WHITE, code_color=GOLD)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def bullets(self, title, items, kind="Lecture", callout=None, notes=""):
        """Full-width bullets; optional callout=(head, body) along the bottom."""
        s = self._slide()
        self._title(s, title, kind)
        bottom = Inches(5.2) if callout else Inches(6.7)
        self._bullets(s, MARGIN, Inches(1.4), W - 2 * MARGIN, bottom - Inches(1.4), items, 20)
        if callout:
            self._callout(s, MARGIN, Inches(5.35), W - 2 * MARGIN, Inches(1.4), callout[0], callout[1])
        s.notes_slide.notes_text_frame.text = notes
        return s

    def bullets_code(self, title, items, code, kind="Lecture", caption=None, notes="", code_size=None,
                     split=0.46):
        """Bullets on the left, a code panel on the right."""
        s = self._slide()
        self._title(s, title, kind)
        lw = (W - 2 * MARGIN) * split
        self._bullets(s, MARGIN, Inches(1.45), lw, Inches(5.3), items, 18)
        cx = MARGIN + lw + Inches(0.35)
        cw = W - MARGIN - cx
        ch = Inches(4.9) if caption else Inches(5.3)
        self._code(s, cx, Inches(1.45), cw, ch, code, code_size)
        if caption:
            self._text(s, cx, Inches(6.45), cw, Inches(0.4), caption, 13, MUTED)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def compare(self, title, left_head, left_items, right_head, right_items, kind="Lecture", notes=""):
        s = self._slide()
        self._title(s, title, kind)
        cw = (W - 2 * MARGIN - Inches(0.4)) / 2
        for i, (head, items, tint) in enumerate(((left_head, left_items, SURFACE),
                                                 (right_head, right_items, RGBColor(0xE8, 0xF0, 0xFD)))):
            x = MARGIN + i * (cw + Inches(0.4))
            self._box(s, x, Inches(1.45), cw, Inches(5.25), tint, radius=0.04)
            self._text(s, x + Inches(0.3), Inches(1.65), cw - Inches(0.6), Inches(0.5), head, 22, PURPLE if i == 0 else LECTURE, bold=True)
            self._bullets(s, x + Inches(0.3), Inches(2.3), cw - Inches(0.6), Inches(4.3), items, 17)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def exercise(self, num, title, minutes, folder, steps, checkin, notes=""):  # checkin = self-check text
        s = self._slide()
        self._title(s, f"Exercise {num}: {title}", "Hands-on")
        # left: meta card
        self._box(s, MARGIN, Inches(1.45), Inches(3.4), Inches(5.25), SURFACE, radius=0.05)
        self._text(s, MARGIN + Inches(0.3), Inches(1.7), Inches(2.8), Inches(0.4), "TIME", 12, MUTED, bold=True)
        self._text(s, MARGIN + Inches(0.3), Inches(2.0), Inches(2.8), Inches(0.8), minutes, 36, HANDSON, bold=True)
        self._text(s, MARGIN + Inches(0.3), Inches(3.05), Inches(2.8), Inches(0.4), "FOLDER", 12, MUTED, bold=True)
        self._text(s, MARGIN + Inches(0.3), Inches(3.35), Inches(2.8), Inches(0.9), f"`{folder}`", 16, TEXT)
        self._text(s, MARGIN + Inches(0.3), Inches(4.35), Inches(2.8), Inches(0.4), "SELF-CHECK", 12, MUTED, bold=True)
        self._text(s, MARGIN + Inches(0.3), Inches(4.65), Inches(2.8), Inches(1.9), checkin, 15, TEXT)
        # right: numbered steps
        x = MARGIN + Inches(3.8)
        y = Inches(1.5)
        rowh = min(Inches(0.95), Inches(5.2) / max(len(steps), 1))
        for i, st in enumerate(steps):
            self._circle_num(s, x, y + Inches(0.02), Inches(0.42), str(i + 1), fill=HANDSON, size=14)
            self._text(s, x + Inches(0.6), y, W - MARGIN - x - Inches(0.6), rowh, st, 17, TEXT)
            y += rowh
        s.notes_slide.notes_text_frame.text = notes
        return s

    def walkthrough(self, title, points, kind="Solutions", code=None, notes=""):
        if code:
            return self.bullets_code(title, points, code, kind=kind, notes=notes)
        return self.bullets(title, points, kind=kind, notes=notes)

    def image(self, title, path, kind="Lecture", caption=None, notes=""):
        """A picture (PNG/JPG), scaled to fit below the title and centered; optional caption."""
        from PIL import Image
        s = self._slide()
        self._title(s, title, kind)
        top = Inches(1.3)
        box_w = W - 2 * MARGIN
        box_h = (Inches(6.35) if caption else Inches(6.85)) - top
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(box_w / iw, box_h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        s.shapes.add_picture(path, int((W - pw) / 2), top, pw, ph)
        if caption:
            self._text(s, MARGIN, top + ph + Inches(0.08), box_w, Inches(0.45), caption, 14, MUTED,
                       align=PP_ALIGN.CENTER)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def selftest(self, num, questions, notes=""):
        s = self._slide()
        self._title(s, f"Self-test: Part {num}")
        y = Inches(1.45)
        rowh = Inches(5.2) / len(questions)
        for i, q in enumerate(questions):
            self._circle_num(s, MARGIN, y + Inches(0.05), Inches(0.5), f"Q{i + 1}", fill=SURFACE, color=PURPLE, size=14)
            self._text(s, MARGIN + Inches(0.75), y + Inches(0.05), W - 2 * MARGIN - Inches(0.75), rowh, q, 18, TEXT)
            y += rowh
        self._text(s, MARGIN, Inches(6.55), W - 2 * MARGIN, Inches(0.35),
                   "Answers are on the lab page: click “Show answer” under each question.", 13, MUTED)
        s.notes_slide.notes_text_frame.text = notes
        if num >= 1:          # every part with an exercise ends with its Canvas quiz (Lab 1 Part 0 has none)
            self.section_quiz(num)
        return s

    def section_quiz(self, num):
        lab = self.lab_label.split(":")[0]           # "Lab 3"
        quiz = f"{lab} Exercise {num} Quiz"
        return self.bullets("Section Quiz on Canvas", [
            f"Open the quiz “{quiz}” on Canvas: your instructor makes it available now",
            "10 minutes, multiple choice: the questions about Exercise %d" % num,
            "The last thing in this part, after the walkthrough and the self-test"],
            kind="Quiz", notes=f"Publish (or unlock) {quiz} in Canvas now; give the class 10 minutes.")

    def submission(self, lab, tree, steps, notes=""):
        s = self._slide()
        self._title(s, "What to submit: one zip, once")
        lw = (W - 2 * MARGIN) * 0.5
        y = Inches(1.5)
        for i, st in enumerate(steps):
            self._circle_num(s, MARGIN, y, Inches(0.42), str(i + 1), size=14)
            self._text(s, MARGIN + Inches(0.6), y - Inches(0.02), lw - Inches(0.6), Inches(1.0), st, 17, TEXT)
            y += Inches(1.0)
        cx = MARGIN + lw + Inches(0.35)
        self._code(s, cx, Inches(1.45), W - MARGIN - cx, Inches(5.3), tree, 14)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def takeaways(self, items, next_lab=None, notes=""):
        s = self._slide(dark=True)
        self._text(s, MARGIN, Inches(0.6), Inches(12), Inches(0.9), "Takeaways", 40, WHITE, bold=True, font=HEAD_FONT)
        n = len(items)
        cols = 2
        cw = (W - 2 * MARGIN - Inches(0.4)) / cols
        rh = Inches(2.35)
        for i, it in enumerate(items):
            x = MARGIN + (i % cols) * (cw + Inches(0.4))
            y = Inches(1.7) + (i // cols) * (rh + Inches(0.3))
            self._box(s, x, y, cw, rh, PURPLE, radius=0.06)
            self._circle_num(s, x + Inches(0.3), y + Inches(0.3), Inches(0.5), str(i + 1), fill=DARK, color=GOLD, size=16)
            tb = self._text(s, x + Inches(1.0), y + Inches(0.25), cw - Inches(1.3), rh - Inches(0.5), "",
                            anchor=MSO_ANCHOR.MIDDLE)
            p = tb.text_frame.paragraphs[0]
            self._runs(p, it, 22, WHITE, code_color=GOLD)
        if next_lab:
            self._text(s, MARGIN, Inches(6.85), Inches(12), Inches(0.4), next_lab, 14, GOLD)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self, path):
        self.prs.save(path)
